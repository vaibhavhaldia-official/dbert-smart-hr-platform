"""
Shared pytest fixtures.

Every test runs against a temporary database and temporary output directories.
No test may touch the developer's gayatri.db, open an SMTP connection, or call
a Google API (AC-14).
"""

import os
import sys
from pathlib import Path

import pytest

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))


# Prevent config.py from loading the developer's .env at import time. Tests must
# never inherit real credentials, and a test that picks one up can make a live
# SMTP connection (AC-14: the suite runs fully offline).
os.environ["GAYATRI_SKIP_DOTENV"] = "1"
for _key in list(os.environ):
    if _key.startswith("GAYATRI_") and _key != "GAYATRI_SKIP_DOTENV":
        del os.environ[_key]


@pytest.fixture(autouse=True)
def _block_network(monkeypatch):
    """Fail loudly if any test tries to open a real socket connection."""
    import socket

    real_connect = socket.socket.connect

    def guard(self, address, *args, **kwargs):
        host = address[0] if isinstance(address, tuple) else address
        if host in ("127.0.0.1", "::1", "localhost"):
            return real_connect(self, address, *args, **kwargs)
        raise AssertionError(
            "Test attempted a network connection to %r. The suite must run "
            "fully offline (AC-14); mock the transport instead." % (address,)
        )

    monkeypatch.setattr(socket.socket, "connect", guard)


@pytest.fixture
def temp_db(tmp_path, monkeypatch):
    """Point the application at a throwaway database and reload the modules."""
    db_path = tmp_path / "test.db"
    monkeypatch.setenv("GAYATRI_DB", str(db_path))

    import importlib
    import config
    importlib.reload(config)
    import bot_config
    importlib.reload(bot_config)
    import database
    importlib.reload(database)

    database.init_database()
    yield database
    importlib.reload(config)
    importlib.reload(bot_config)
    importlib.reload(database)


@pytest.fixture
def sample_docx(tmp_path):
    """Build a minimal DOCX with placeholders, so no binary asset is committed."""
    from docx import Document

    doc = Document()
    doc.add_paragraph("Offer Letter")
    doc.add_paragraph("Dear {name},")
    doc.add_paragraph("Your email is {email}.")
    doc.add_paragraph("Domain: {internship}")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Start"
    table.cell(0, 1).text = "{start_date}"
    table.cell(1, 0).text = "Stipend"
    table.cell(1, 1).text = "{stipend_total}"
    path = tmp_path / "sample_offer.docx"
    doc.save(str(path))
    return path


class FakeSMTP:
    """In-memory stand-in for smtplib — captures messages, sends nothing."""

    sent = []

    def __init__(self, *args, **kwargs):
        self.messages = []

    def __enter__(self):
        return self

    def __exit__(self, *args):
        return False

    def starttls(self, *args, **kwargs):
        pass

    def login(self, *args, **kwargs):
        pass

    def send_message(self, msg, *args, **kwargs):
        FakeSMTP.sent.append(msg)

    def sendmail(self, from_addr, to_addrs, msg, *args, **kwargs):
        FakeSMTP.sent.append(msg)

    def quit(self):
        pass


@pytest.fixture
def fake_smtp(monkeypatch):
    """Replace smtplib transports so no test can reach a mail server."""
    import smtplib

    FakeSMTP.sent = []
    monkeypatch.setattr(smtplib, "SMTP", FakeSMTP)
    monkeypatch.setattr(smtplib, "SMTP_SSL", FakeSMTP)
    return FakeSMTP
